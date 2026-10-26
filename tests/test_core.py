import io
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from fuel.core import Settings, analyze, detect_columns, normalize_plate, plate_typos, read_any, read_rates, summary, to_date, to_number
from fuel.report import build_excel


def frame(rows, cols=("Гос. номер", "Водитель", "Дата", "Пробег, км", "Расход факт, л", "Норма, л/100 км")):
    return pd.DataFrame(rows, columns=list(cols))


def test_column_detection_ru_kk_en():
    assert detect_columns(frame([]))["norm_rate"] == "Норма, л/100 км"
    m = detect_columns(pd.DataFrame(columns=["Мемлекеттік нөмір", "Жүргізуші", "Жүріс", "Нақты шығын", "Норма, л"]))
    assert m == {"plate": "Мемлекеттік нөмір", "driver": "Жүргізуші", "mileage": "Жүріс", "fact": "Нақты шығын", "norm_l": "Норма, л"}
    m = detect_columns(pd.DataFrame(columns=["Vehicle", "Distance", "Fuel used", "Norm liters"]))
    assert {"plate", "mileage", "fact", "norm_l"} <= set(m)


def test_value_cleaning():
    assert normalize_plate(" 101 АВС 01 ") == "101ABC01"          # кириллица → латиница
    assert to_number("38,5") == 38.5 and to_number("1 250") == 1250 and to_number("сорок") is None
    assert to_date("14.10.2026") == pd.Timestamp("2026-10-14")
    assert to_date("2026-10-14") == pd.Timestamp("2026-10-14")
    assert to_date("31.09.2026") is None
    assert to_date(46309) == pd.Timestamp("2026-10-14")             # дата Excel числом


def test_deviation_math_and_flags():
    det, issues, tot = analyze(frame([
        ["101ABC01", "А", "01.10.2026", 100, 16, 16],    # ровно норма
        ["101ABC01", "А", "02.10.2026", 200, 40, 16],    # норма 32, +8 л, +25% → проблема
        ["101ABC01", "А", "03.10.2026", 50, 9, 16],      # норма 8, +1 л, +12.5% — меньше 5 л, не проблема
        ["101ABC01", "А", "04.10.2026", 0, 30, 16],      # расход без пробега
    ]))
    assert list(det["Норма, л"]) == [16, 32, 8, 0]
    assert list(det["Отклонение, л"]) == [0, 8, 1, 30]
    assert det["Отклонение, %"].iloc[1] == 25.0
    assert list(det["Проблема"]) == [False, True, False, True]
    assert "без пробега" in det["Признак"].iloc[3]
    assert tot["problems"] == 2 and tot["vehicles"] == 1 and not issues


def test_input_errors_reported_not_crashing():
    det, issues, tot = analyze(frame([
        ["101ABC01", "А", "01.10.2026", 100, 16, 16],
        ["", "Б", "01.10.2026", 100, 16, 16],
        ["101ABC01", "А", "31.09.2026", 100, 16, 16],
        ["101ABC01", "А", "02.10.2026", None, 16, 16],
        ["101ABC01", "А", "03.10.2026", 100, "сорок", 16],
        ["101ABC01", "А", "04.10.2026", -5, 16, 16],
    ]))
    msgs = [(i.row, i.level, i.message) for i in issues]
    assert (3, "ошибка", "пустой госномер — строка не учтена") in msgs
    assert any(r == 4 and "даты" in m for r, _, m in msgs)
    assert any(r == 5 and "пустая ячейка" in m for r, _, m in msgs)
    assert any(r == 6 and "сорок" in m for r, _, m in msgs)
    assert any(r == 7 and "отрицательное" in m for r, _, m in msgs)
    assert tot["rows_ok"] == 2


def test_plate_typo_merged():
    plates = ["214KTE01"] * 5 + ["214KTF01", "305MNA01", "305MNA01"]
    assert plate_typos(plates) == {"214KTF01": "214KTE01"}
    det, issues, _ = analyze(frame([[p, "", "01.10.2026", 100, 16, 16] for p in plates]))
    assert set(det["Госномер"]) == {"214KTE01", "305MNA01"}
    assert any("опечатка" in i.message for i in issues)


def test_norm_from_reference_and_winter():
    rates = read_rates(pd.DataFrame({"Модель": ["КамАЗ 65115"], "Норма, л/100 км": [34]}))
    df = frame([["305MNA01", "КамАЗ 65115", "01.12.2026", 100, 40]],
               cols=("Госномер", "Модель", "Дата", "Пробег, км", "Факт, л"))
    det, _, _ = analyze(df, Settings(rates=rates, winter_coef_pct=10))
    assert det["Норма, л"].iloc[0] == pytest.approx(37.4)            # 34 + 10% зимой


def test_missing_columns_message():
    with pytest.raises(ValueError, match="пробег"):
        analyze(pd.DataFrame({"Госномер": ["1"], "Факт": [1], "Норма, л": [1]}))


def test_csv_semicolon_and_excel_report():
    csv = "Госномер;Водитель;Пробег, км;Факт, л;Норма, л\n101ABC01;Ахметов;100;20,5;16\n".encode("utf-8")
    df = read_any(csv, "a.csv")
    det, issues, tot = analyze(df)
    assert det["Отклонение, л"].iloc[0] == 4.5
    wb = load_workbook(io.BytesIO(build_excel(det, issues, tot, Settings(), "a.csv")))
    assert wb.sheetnames == ["Сводка", "По машинам", "По водителям", "Проблемы", "Детально", "Ошибки в данных"]
    assert wb["По машинам"]["A2"].value == "101ABC01"


def test_sample_file_end_to_end():
    data = open("samples/primer_avtopark_oktyabr.xlsx", "rb").read()
    det, issues, tot = analyze(read_any(data, "x.xlsx"))
    top = summary(det, "Госномер").iloc[0]
    assert top["Госномер"] == "305MNA01" and top["Отклонение, %"] > 15     # заложенный перерасход найден
    assert tot["errors"] == 3 and tot["warnings"] == 2
