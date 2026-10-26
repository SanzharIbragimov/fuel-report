"""Итоговый отчёт в Excel: сводка, машины, водители, проблемы, детально, предупреждения."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core import Issue, Settings, summary

HEAD = PatternFill("solid", fgColor="1F4E78")
RED = PatternFill("solid", fgColor="F8D7DA")
AMBER = PatternFill("solid", fgColor="FFF3CD")
GREEN = PatternFill("solid", fgColor="D1E7DD")
THIN = Side(style="thin", color="D0D7DE")


def _sheet(ws, df: pd.DataFrame, widths: dict | None = None, pct_cols=(), red_rows=None, limit: float = 10.0):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
                   for v in row])
    n = len(df) + 1
    for j, col in enumerate(df.columns, 1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = (widths or {}).get(col, max(12, min(45, len(str(col)) + 4)))
        for cell in ws[letter][1:]:
            cell.border = Border(bottom=THIN)
            if col == "Дата":
                cell.number_format = "DD.MM.YYYY"
            elif col in pct_cols:
                cell.number_format = '+0.0"%";-0.0"%";0"%"'
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.0"
        if col in pct_cols and n > 1:      # цвет по отклонению
            rng = f"{letter}2:{letter}{n}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=[str(limit)], fill=RED))
            ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.0001", str(limit)], fill=AMBER))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=["0"], fill=GREEN))
    if red_rows is not None:
        for i, bad in enumerate(red_rows, start=2):
            if bad:
                for cell in ws[i]:
                    cell.fill = RED
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"          # печать: по ширине листа
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if n > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{n}"


def build_excel(det: pd.DataFrame, issues: list[Issue], totals: dict, st: Settings, source_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Отчёт по расходу топлива"
    ws["A1"].font = Font(bold=True, size=14)
    period = ""
    if len(det) and det["Дата"].notna().any():
        period = f"{det['Дата'].min():%d.%m.%Y} — {det['Дата'].max():%d.%m.%Y}"
    lines = [
        ("Файл", source_name), ("Период", period or "не указан"),
        ("Сформирован", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Порог перерасхода", f"> {st.threshold_pct:g}% и ≥ {st.threshold_l:g} л"),
        ("Зимняя надбавка к норме", f"{st.winter_coef_pct:g}%" if st.winter_coef_pct else "не применялась"),
        (None, None),
        ("Машин", totals["vehicles"]), ("Записей учтено", f"{totals['rows_ok']} из {totals['rows_in']}"),
        ("Расход по факту, л", round(totals["fact"], 1)), ("Расход по норме, л", round(totals["norm"], 1)),
        ("Отклонение по автопарку", f"{totals['dev_pct']:+.1f}%" if totals["dev_pct"] is not None else "—"),
        ("Перерасход (сумма превышений), л", round(totals["overrun_l"], 1)),
        ("Проблемных записей", totals["problems"]),
        ("Ошибок во входных данных", totals["errors"]), ("Предупреждений", totals["warnings"]),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        ws.cell(i, 1, k).font = Font(bold=True) if k else Font()
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 36, 34
    if totals["problems"]:
        ws.cell(len(lines) + 4, 1, "Проблемные записи — на листе «Проблемы», машины с перерасходом выделены на листе «По машинам».")

    by_car = summary(det, "Госномер")
    _sheet(wb.create_sheet("По машинам"), by_car, {"Госномер": 14, "Модель": 20}, pct_cols=("Отклонение, %",), limit=st.threshold_pct)
    by_driver = summary(det, "Водитель")
    if len(by_driver):
        _sheet(wb.create_sheet("По водителям"), by_driver, {"Водитель": 28}, pct_cols=("Отклонение, %",), limit=st.threshold_pct)
    cols = ["Строка", "Дата", "Госномер", "Модель", "Водитель", "Пробег, км", "Факт, л", "Норма, л",
            "Отклонение, л", "Отклонение, %", "Признак"]
    prob = det[det["Проблема"]].sort_values("Отклонение, %", ascending=False)[cols]
    _sheet(wb.create_sheet("Проблемы"), prob, {"Водитель": 26, "Признак": 55, "Модель": 18},
           pct_cols=("Отклонение, %",), limit=st.threshold_pct)
    _sheet(wb.create_sheet("Детально"), det[cols], {"Водитель": 26, "Признак": 55, "Модель": 18},
           pct_cols=("Отклонение, %",), red_rows=det["Проблема"].tolist(), limit=st.threshold_pct)
    iss = pd.DataFrame([(x.row, x.level, x.column, x.message) for x in sorted(issues, key=lambda x: x.row)],
                       columns=["Строка файла", "Уровень", "Поле", "Что не так"])
    _sheet(wb.create_sheet("Ошибки в данных"), iss, {"Что не так": 60}, red_rows=(iss["Уровень"] == "ошибка").tolist())
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
